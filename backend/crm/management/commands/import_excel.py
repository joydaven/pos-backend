import os
from datetime import datetime
import json   # 🔧 added

import openpyxl
import re
from django.core.management.base import BaseCommand
from django.utils import timezone

from crm.models import WooServiceTypes  # adjust if your model is elsewhere


class Command(BaseCommand):
    help = "Import WooServiceTypes from Excel file"

    def add_arguments(self, parser):
        parser.add_argument("excel_file", type=str, help="Path to Excel file")
        parser.add_argument(
            "--sheet",
            type=str,
            default="Credited Services",
            help="Name of the sheet to import (default: 'Credited Services')",
        )

    def handle(self, *args, **options):
        excel_file = options["excel_file"]
        sheet_name = options["sheet"]

        if not os.path.exists(excel_file):
            self.stderr.write(self.style.ERROR(f"File not found: {excel_file}"))
            return

        wb = openpyxl.load_workbook(excel_file, data_only=True)

        # Select the target sheet
        try:
            sheet = wb[sheet_name]
        except KeyError:
            self.stderr.write(self.style.ERROR(f"Sheet '{sheet_name}' not found in workbook"))
            return

        self.stdout.write(f"Using sheet: {sheet.title}")

        # Extract headers
        headers = [str(c).strip() if c else "" for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        normalized = [h.lower().replace(" ", "").replace("\n", "") for h in headers]

        self.stdout.write(f"DEBUG HEADERS: {headers}")

        try:
            name_idx = normalized.index("productname")
            series_idx = normalized.index("series")
        except ValueError as e:
            self.stderr.write(self.style.ERROR(f"Missing expected column: {e}"))
            return

        imported, updated = 0, 0

        # Iterate rows
        for row in sheet.iter_rows(min_row=2, values_only=True):
            name = row[name_idx]
            series_raw = row[series_idx] if series_idx < len(row) else None

            # Skip empty rows
            if not name or str(name).strip() == "":
                continue

            series_list = self.parse_series(series_raw)

            # 🔧 Ensure JSON-safe value
            try:
                json.dumps(series_list)  # will raise if not serializable
            except TypeError:
                series_list = [str(x) for x in series_list]

            self.stdout.write(
                f"DEBUG: {name} | raw={series_raw!r} | parsed={series_list}"
            )

            obj, created = WooServiceTypes.objects.update_or_create(
                name=name,
                defaults={
                    "series": series_list,  # ✅ force JSON
                    "updated_at": timezone.now(),
                },
            )

            if created:
                imported += 1
            else:
                updated += 1

        self.stdout.write(
            self.style.SUCCESS(f"Import finished: {imported} new, {updated} updated.")
        )

    def parse_series(self, raw):
        if raw is None:
            return []

        # If Excel turned "6/2/12" into a date, extract all components
        if isinstance(raw, datetime):
            return [raw.month, raw.day, raw.year % 100]  # e.g. 2012-06-03 -> [6, 3, 12]

        if isinstance(raw, (int, float)):
            return [int(raw)]

        if isinstance(raw, str):
            import re
            series_list = []
            parts = re.split(r"[,\s;]+", raw)
            for part in parts:
                if part.isdigit():
                    series_list.append(int(part))
            return series_list

        return []
